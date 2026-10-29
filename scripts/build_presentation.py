"""Build the review deck from archived figures. Requires python-pptx>=1.0,<2."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'presentation'
FIG = ROOT / 'docs' / 'research_results'
NAVY, TEAL, GRAY, WHITE = '142D44', '087F8C', '526474', 'FFFFFF'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'NeuroText: EEG-to-sentence retrieval'
prs.core_properties.author = 'Ziad Elsayed'
prs.core_properties.subject = 'Draft review; cross-dataset experiments pending identity evidence'


def box(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, content, size=22, color=NAVY, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for i, line in enumerate(content.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
    return shape


def slide(title, subtitle, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    box(s, 0, 0, 13.333, .12, TEAL)
    text(s, .55, .35, 12.2, .65, title, 30, bold=True)
    text(s, .55, 1.05, 12.2, .62, subtitle, 16, GRAY)
    text(s, .55, 7.1, 11.6, .25, 'NEUROTEXT  /  REVIEW DRAFT  /  29 OCT 2026', 10, GRAY)
    text(s, 12.1, 7.1, .7, .25, f'{len(prs.slides):02}', 10, GRAY)
    s.notes_slide.notes_text_frame.text = notes
    return s


def callout(s, message, y=6.32):
    box(s, .55, y, 12.2, .58, 'E8F5F4')
    text(s, .73, y + .1, 11.8, .4, message, 17, TEAL, True)


def picture(s, filename, x=.65, y=1.75, w=12, h=4.32):
    path = FIG / filename
    with Image.open(path) as im:
        ratio = im.width / im.height
    width, height = min(w, h * ratio), min(h, w / ratio)
    s.shapes.add_picture(str(path), Inches(x + (w-width)/2), Inches(y + (h-height)/2),
                         width=Inches(width), height=Inches(height))


def chart(title, subtitle, filename, takeaway, notes):
    s = slide(title, subtitle, notes + f'\nFigure: docs/research_results/{filename}')
    picture(s, filename)
    callout(s, takeaway)


s = slide('NeuroText', 'EEG representation learning for sentence retrieval',
          'Present as a draft based on Windows RTX 3050 execution. This deck is not a completed '
          'cross-dataset study. Source: docs/research_results/REPORT.md and retrieval.csv (79 reports).')
text(s, .7, 2, 11.8, 1.5, 'Can EEG retrieve sentence meaning\nbeyond matched negative controls?', 36, bold=True)
text(s, .7, 4, 11.8, 1.3, 'ZuCo 1.0 + 2.0  |  Statistical ridge · CNN · SSL · frozen LaBraM\nZiad Elsayed  |  University project', 22, GRAY)
callout(s, 'Executed within-dataset results; strict cross-dataset evaluation remains pending.')

s = slide('Define the prediction task', 'Fixed-pool retrieval with unseen readers',
          'Source: REPORT.md, Data and methods. The text encoder is frozen MiniLM. MRR averages '
          'reciprocal rank of the correct sentence. R@K is the fraction with the correct target '
          'in the first K positions. Pool size changes expected random performance.')
text(s, .75, 1.95, 11.8, 3.8,
     '1. Encode a sentence-length EEG recording.\n'
     '2. Predict its location in a frozen sentence-embedding space.\n'
     '3. Rank unique candidate sentences by cosine similarity.\n'
     '4. Evaluate Recall@1/5/10 and mean reciprocal rank (MRR).', 25)
callout(s, 'Above random ranking alone is insufficient: real EEG must beat matched controls.')

s = slide('Data preparation and quality audit', 'Normal reading · 105 ordered channels · resampled to 200 Hz',
          'Source: REPORT.md. ZuCo1: 3356 prepared, 7 duration-QC exclusions, 3349 retained. '
          'ZuCo2: 5911 prepared, 27 exclusions, 5884 retained. Missing source slots were separately '
          'audited. Original timing/reference/filter histories remain incompletely verified. '
          'LaBraM adaptation has different counts: 3345 and 5838, requiring whole seconds.')
for x, name, people, prepared, retained in [(.75, 'ZuCo 1.0', '12 readers', '3,356', '3,349'),
                                           (6.85, 'ZuCo 2.0', '18 readers', '5,911', '5,884')]:
    box(s, x, 1.95, 5.7, 3.5, 'F1F5F8')
    text(s, x+.25, 2.2, 5.1, .55, name, 28, TEAL, True)
    text(s, x+.25, 2.95, 5.1, 2.2,
         f'{people}\n{prepared} prepared observations\n{retained} after duration QC', 24)
callout(s, '19-channel LaBraM adaptation is a separate cohort: 3,345 / 5,838 observations.')

s = slide('One shared, memory-conscious implementation', 'Windows + Colab clients use the same Python package',
          'Source: REPORT.md and project source modules. Colab compatibility is implemented; '
          'actual Colab execution has not been established. Statistical features include per-channel '
          'statistics and spectral power. SSL reconstructs masked temporal blocks before alignment. '
          'Official LaBraM is frozen and verified, not fine-tuned. Local hardware: RTX 3050, 4 GB VRAM.')
for x, heading, body in [(.7, 'LOAD', 'Lazy signal arrays\nAudited manifests'),
                          (3.9, 'REPRESENT', 'Statistics / CNN\nSSL / frozen LaBraM'),
                          (7.1, 'ALIGN', 'Train-only scaling\nValidation selection'),
                          (10.3, 'RANK', 'MiniLM targets\nCosine retrieval')]:
    box(s, x, 2.25, 2.35, 2.3, 'F1F5F8')
    text(s, x+.15, 2.5, 2.05, .5, heading, 19, TEAL, True)
    text(s, x+.15, 3.15, 2.05, 1.3, body, 18)
text(s, .8, 5, 11.8, .8, 'Local hardware: RTX 3050 with 4 GB VRAM; lazy loading limits memory use.', 22)
callout(s, 'Shared code and documented inputs make weak results inspectable and reproducible.')

s = slide('Controls make the central comparison', 'Hold out subjects; fit preprocessing on training observations only',
          'Source: REPORT.md. Gaussian noise preserves each observation’s channel standard deviations, '
          'so it retains amplitude information and is not uniform random ranking. Shuffled training '
          'labels are deranged by sentence class; evaluation labels are true. For SSL, Gaussian covers '
          'pretraining and alignment. Shuffled alignment may reuse label-free real-data pretraining. '
          'Test inspection during development means this is exploratory research.')
text(s, .8, 1.95, 11.8, 3.8,
     'Real EEG: original signals and true sentence pairing.\n'
     'Gaussian: matched per-channel scale, independently generated noise.\n'
     'Shuffled: separately trained model with incorrect sentence pairing.\n'
     'Compare within the same cohort, split, seed and candidate pool.', 25)
callout(s, 'Single-split differences are descriptive; no statistical significance is claimed.')

chart('Supervised CNN: Gaussian matches top-1 hits',
      'Seed 42 · ZuCo1: 592 queries / 300 candidates · ZuCo2: 984 / 344',
      'cnn_controls_mrr.png', 'Gaussian MRR exceeds real EEG in both datasets.',
      'Source: REPORT.md, Supervised CNN negative controls. ZuCo1 uses original cohort and an '
      '8-second center crop; ZuCo2 uses duration QC and full recordings. These are not a controlled '
      'cross-dataset comparison. Real and Gaussian top-1 hits: 2 in ZuCo1, 3 in ZuCo2. '
      'Dashed lines are uniform random-ranking expectations, not significance thresholds.')

chart('SSL: reconstruction does not establish decoding',
      'Full recordings · duration QC · Gaussian control covers both learned stages',
      'ssl_controls_mrr.png', 'No consistent retrieval advantage for real EEG over the controls.',
      'Source: REPORT.md, Full-sentence SSL negative controls. Query/candidate counts: 592/300 '
      'and 984/344. Real/Gaussian/shuffled MRR: ZuCo1 .020680/.023159/.019842; '
      'ZuCo2 .018846/.018117/.019402. Training improvement does not imply generalization. '
      'One split; no significance claim.')

chart('Frozen LaBraM against matched statistical features',
      'Identical adapted signals · 19 mapped electrodes · whole-second support',
      'pretrained_comparison.png', 'Statistical features have higher top-1 and MRR on both reported splits.',
      'Source: REPORT.md, Matched pretrained comparison. Official weights checksum verified. '
      'Template-based mapping errors 2.11–9.40 degrees; approximate mapping and unknown upstream '
      'reference/filter details limit general conclusions. ZuCo1 592/300; ZuCo2 983/344. '
      'ZuCo1 LaBraM Recall@5/10 is higher; do not claim all metrics favor statistics.')

chart('LaBraM controls limit the interpretation',
      'Frozen backbone · separately fit ridge models · same adapted cohort and split',
      'labram_controls_mrr.png', 'No consistent real-EEG superiority under this exploratory adaptation.',
      'Source: REPORT.md, Frozen LaBraM ridge negative controls. Gaussian MRR exceeds real in '
      'ZuCo1; ZuCo2 real MRR is slightly higher but top-1 has fewer hits (4 versus 5). '
      'Duplicate report sets are displayed once when all exported fields agree; this is not '
      'independent replication or a general verdict on LaBraM.')

chart('Unseen subjects and unseen sentence content',
      'Ridge · seeds 42, 43, 44 · 60 candidates in ZuCo1 / 68 in ZuCo2',
      'unseen_content_mrr.png', 'Real EEG exceeds Gaussian MRR on only one of three seeds in each dataset.',
      'Source: REPORT.md, Unseen subjects and unseen sentence content. Points show individual '
      'splits; report SD across splits is not a confidence interval. Splits reuse subjects and stimuli. '
      'Real/Gaussian mean MRR: .085268/.086335 in ZuCo1; .073448/.075441 in ZuCo2. '
      'Do not compare raw recall percentages against 300/344-candidate experiments.')

chart('Quality weighting shows no consistent benefit',
      'ZuCo2 ridge · 984 queries / 344 candidates · one split and corruption seed',
      'quality_robustness_mrr.png', 'Quality-aware MRR is lower in 8 of 9 nonzero corruption conditions.',
      'Source: REPORT.md, Robustness and quality weighting. Clean MRR: standard .023036, '
      'quality-aware .020230. Exception: 25% channel dropout. Nonmonotonic curves near chance '
      'do not establish beneficial signal removal. Noise, channel dropout and temporal masking '
      'have different severity units. No broad robustness claim.')

s = slide('Published work: compare protocols before scores', 'Generation, classification and retrieval measure different tasks',
          'Sources: docs/published_work_comparison.md; docs/research_results/literature_sources.json. '
          'Wang and Ji, AAAI 2022: Open Vocabulary EEG-to-Text Decoding and Zero-shot Sentiment '
          'Classification. Official repository acknowledges evaluation correction in 2024 and '
          'directs readers to NeuSpeech. LaBraM: Jiang, Zhao and Lu, ICLR 2024. '
          'Review pins primary source revisions. This project has no protocol-matched numerical '
          'paper reproduction. Other limits: small cohorts, montage/reference uncertainty and '
          'repeated test inspection during development.')
text(s, .8, 1.95, 11.8, 3.8,
     'EEG-to-text papers often report BLEU / ROUGE for language generation.\n'
     'LaBraM originally reports classification tasks.\n'
     'Our task reports Recall@K / MRR for a fixed sentence pool.\n'
     'The original AAAI 2022 repository documents an evaluation correction.', 24)
callout(s, 'Primary-source protocol review is complete; competitive numerical ranking is unsupported.')

s = slide('Cross-dataset transfer: prepared, not executed', 'Both directions remain pending participant identity evidence',
          'Source: user execution transcript, transfer_plan_20261029_213051/plan.json summary. '
          '9233 observations; 594 unique targets; text cache rebuilt at MiniLM revision '
          '1110a243fdf4706b3f48f1d95db1a4f5529b4d41. The summary has been supplied, '
          'not the full plan JSON. No EEG copies or training. Maintainer inquiry was sent by '
          'the user; no response supplied. Different dataset ID prefixes do not establish '
          'participant independence. Existing selected NR manifests have 50 shared canonical texts; '
          'transfer must disclose overlap. Original signals and completed results are preserved.')
text(s, .8, 1.9, 11.8, 3.8,
     'Prepared: 9,233 signal references and 594 shared sentence targets.\n'
     'Directions: ZuCo1 → ZuCo2 and ZuCo2 → ZuCo1.\n'
     'Pending: verify whether the same people appear in both releases.\n'
     'Then: reconcile identities and fit / select models on the source only.', 25)
callout(s, 'Preparation is successful; it is not an executed transfer result.')

s = slide('Conclusion and remaining work', 'A reproducible negative result with explicit limits',
          'Sources: REPORT.md; archived retrieval.csv, 79 executed evaluation reports including '
          'duplicates and robustness conditions, not 79 independent methods. Software validation '
          '121 passed; synthetic tests verify software and do not prove scientific accuracy. '
          'Remaining: independent identity evidence, strict transfer in both directions, '
          'original JSON/config/checkpoint audit and final submission review.')
text(s, .8, 1.95, 11.8, 3.8,
     'Finding: tested representations do not show reliable retrieval beyond controls.\n'
     'Contribution: modular pipeline, audited data, matched controls and clear limits.\n'
     'Remaining: identity evidence, strict transfer experiments and final review.\n'
     'Archive: original evaluation JSON, configs, splits and checkpoints.', 24)
callout(s, 'These findings do not establish that EEG contains no semantic information.')

OUT.mkdir(parents=True, exist_ok=True)
path = OUT / 'NeuroText_draft.pptx'
prs.save(path)
print(f'Saved {len(prs.slides)} slides: {path}')
